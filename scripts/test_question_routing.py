import shutil
import tempfile
import unittest
from io import BytesIO

from docx import Document
from docx.shared import Inches
from PIL import Image as PILImage

from app.question_routing import detect_sources
from app.states.paper_data import Passage, extract_uploaded_paper
from app.states.workspace_state import (
    _ground_sources,
    _retrieve_passages,
    build_answer_context,
    route_question,
)


def figure() -> bytes:
    stream = BytesIO()
    PILImage.new("RGB", (600, 400), "white").save(stream, format="PNG")
    return stream.getvalue()


def sample_paper():
    document = Document()
    document.add_heading("Abstract", level=1)
    document.add_paragraph(
        "Sixty students recorded outdoor minutes each evening for fourteen days."
    )
    document.add_heading("Methods", level=1)
    document.add_paragraph(
        "Wellbeing was rated from 1 to 7. Outdoor minutes were self-reported."
    )
    document.add_heading("Results", level=1)
    document.add_paragraph("Higher-exposure days scored higher (Hartig et al., 2014).")
    document.add_paragraph("Figure 1. Daily wellbeing ratings by exposure condition.")
    document.add_picture(BytesIO(figure()), width=Inches(4))
    document.add_heading("References", level=1)
    document.add_paragraph(
        "Hartig, T., Mitchell, R., & Frumkin, H. (2014). Nature and health. "
        "Annual Review of Public Health, 35, 207-228."
    )
    document.add_paragraph(
        "Bratman, G. N., Hamilton, J. P., & Gross, J. J. (2015). Nature "
        "experience. PNAS, 112(28), 8567-8572."
    )
    stream = BytesIO()
    document.save(stream)
    return extract_uploaded_paper("green.docx", stream.getvalue())


class QuestionRoutingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.media = tempfile.mkdtemp(prefix="paperproof-rt-")
        import os

        cls._previous = os.environ.get("PAPERPROOF_MEDIA_DIR")
        os.environ["PAPERPROOF_MEDIA_DIR"] = cls.media
        cls.paper = sample_paper()

    @classmethod
    def tearDownClass(cls):
        import os

        shutil.rmtree(cls.media, ignore_errors=True)
        if cls._previous is None:
            os.environ.pop("PAPERPROOF_MEDIA_DIR", None)
        else:
            os.environ["PAPERPROOF_MEDIA_DIR"] = cls._previous

    def context_kinds(self, question):
        intent = route_question(question)
        retrieved = _retrieve_passages(self.paper["passages"], question)
        context = build_answer_context(
            self.paper, intent, retrieved, question=question
        )
        return intent, [item["kind"] for item in context], context

    # ------------------------------------------------------------- routing

    def test_paper_questions_never_receive_the_bibliography(self):
        for question in (
            "What is the main finding?",
            "How was the study designed?",
            "What are the limitations?",
            "What did the authors conclude?",
        ):
            with self.subTest(question=question):
                intent, kinds, _ = self.context_kinds(question)
                self.assertFalse(intent["references"])
                self.assertNotIn("reference", kinds)

    def test_reference_questions_receive_references(self):
        for question in (
            "Which references are cited in the text?",
            "Are there references that are never cited?",
            "Does the reference list include a DOI?",
        ):
            with self.subTest(question=question):
                intent, kinds, context = self.context_kinds(question)
                self.assertTrue(intent["references"])
                self.assertIn("reference", kinds)
                self.assertTrue(context)

    def test_citation_use_questions_also_get_the_body_text(self):
        """A reference question about citation use needs both sources."""
        for question in (
            "Which references are cited in the text?",
            "Are there references that are never cited?",
        ):
            with self.subTest(question=question):
                intent, kinds, context = self.context_kinds(question)
                self.assertTrue(intent["references"])
                self.assertIn("reference", kinds)
                self.assertIn("passage", kinds)
                passage_ids = {
                    item["id"] for item in context if item["kind"] == "passage"
                }
                self.assertTrue(passage_ids)

    def test_metadata_reference_questions_do_not_pull_in_the_body(self):
        question = "Does the reference list include a DOI?"
        intent, kinds, context = self.context_kinds(question)
        self.assertTrue(intent["references"])
        self.assertNotIn("passage", kinds)
        self.assertTrue(
            all(item["kind"] == "reference" for item in context)
        )

    def test_figure_questions_receive_the_figure(self):
        for question in ("What does Figure 1 show?", "What is in the chart?"):
            with self.subTest(question=question):
                intent, kinds, _ = self.context_kinds(question)
                self.assertTrue(intent["images"])
                self.assertIn("image", kinds)
                self.assertFalse(intent["references"])

    def test_mixed_question_combines_sources(self):
        intent, kinds, _ = self.context_kinds(
            "Summarize the figure and the uncited references together"
        )
        self.assertTrue(intent["mixed"])
        self.assertIn("reference", kinds)
        self.assertIn("image", kinds)

    def test_external_context_is_detected_but_never_supplied(self):
        intent, kinds, context = self.context_kinds(
            "According to current evidence, is this still true?"
        )
        self.assertTrue(intent["external"])
        self.assertEqual(context, [])

    def test_figure_context_states_that_pixels_were_not_read(self):
        _, _, context = self.context_kinds("What does Figure 1 show?")
        images = [item for item in context if item["kind"] == "image"]
        self.assertTrue(images)
        self.assertIn("not interpreted", images[0]["label"])

    def test_every_context_block_carries_an_id_and_label(self):
        for question in (
            "What is the main finding?",
            "Which references are cited?",
            "What does Figure 1 show?",
        ):
            with self.subTest(question=question):
                _, _, context = self.context_kinds(question)
                for item in context:
                    self.assertTrue(item["id"])
                    self.assertTrue(item["label"])
                    self.assertIn(item["kind"], {"passage", "reference", "image"})

    # ------------------------------------------------------------ grounding

    def test_reference_only_answer_can_be_grounded(self):
        question = "Which references are cited in the text?"
        payload = {
            "answer": "Two bibliography entries are listed in this paper.",
            "evidence_ids": ["r1", "r2"],
            "confidence": "high",
            "limitations": [],
            "unsupported_parts": [],
        }
        answer, evidence = _ground_sources(question, [], payload["answer"], payload["evidence_ids"], self.paper)
        self.assertEqual(answer, payload["answer"])
        self.assertEqual(len(evidence), 2)
        self.assertTrue(all(item["kind"] == "reference" for item in evidence))
        self.assertTrue(all(item["passage_id"] == "" for item in evidence))

    def test_image_only_answer_can_be_grounded(self):
        payload_ids = ["img-01"]
        answer, evidence = _ground_sources(
            "What does Figure 1 show?",
            [],
            "The figure caption describes daily wellbeing ratings.",
            payload_ids,
            self.paper,
        )
        self.assertEqual(len(evidence), 1)
        self.assertEqual(evidence[0]["kind"], "image")
        self.assertIn("only", evidence[0]["detail"])

    def test_unknown_citation_id_rejects_the_whole_answer(self):
        for ids in (["r99"], ["img-99"], ["r1", "made-up-id"], ["nonsense"]):
            with self.subTest(ids=ids):
                answer, evidence = _ground_sources(
                    "Which references are cited?",
                    [],
                    "A composed answer.",
                    ids,
                    self.paper,
                )
                self.assertEqual(evidence, [])
                self.assertIn("does not establish an answer", answer)

    def test_mixed_evidence_is_combined(self):
        # Build the passage context directly so this test checks grounding
        # rather than the retrieval heuristics.
        retrieved = [
            {
                "id": "p1",
                "passage_id": "p1",
                "section": "Results",
                "page": 3,
                "text": "Higher-exposure days scored higher (Hartig et al., 2014).",
            }
        ]
        answer, evidence = _ground_sources(
            "What did the study find?",
            retrieved,
            "Higher exposure days were associated with higher scores.",
            ["p1", "r1"],
            self.paper,
        )
        self.assertEqual(
            answer, "Higher exposure days were associated with higher scores."
        )
        kinds = {item["kind"] for item in evidence}
        self.assertEqual(kinds, {"passage", "reference"})
        self.assertEqual(len(evidence), 2)

    def test_paper_without_figures_or_references_grounds_passages_only(self):
        paper = {
            "passages": [
                Passage(
                    id="p1", section="Abstract", page=1, text="A short abstract."
                )
            ],
            "references": [],
            "images": [],
        }
        answer, evidence = _ground_sources(
            "What is this?", [dict(paper["passages"][0], passage_id="p1")], "A summary.", ["p1"], paper
        )
        self.assertEqual(len(evidence), 1)
        self.assertEqual(evidence[0]["kind"], "passage")


class DetectSourcesTests(unittest.TestCase):
    def test_plain_questions_select_nothing_special(self):
        for question in (
            "What is the main finding?",
            "How was the study designed?",
            "What are the limitations?",
        ):
            with self.subTest(question=question):
                detected = detect_sources(question)
                self.assertFalse(detected["references"])
                self.assertFalse(detected["images"])
                self.assertFalse(detected["external"])

    def test_external_request_needs_an_explicit_phrase(self):
        self.assertTrue(detect_sources("What do experts say about this?")["external"])
        self.assertTrue(detect_sources("Compare with recent research")["external"])
        self.assertFalse(
            detect_sources("What do the authors say about this?")["external"]
        )


if __name__ == "__main__":
    unittest.main()