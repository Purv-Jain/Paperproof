import reflex as rx

from io import BytesIO
import unittest

from docx import Document
from pypdf import PdfWriter
import pymupdf
from docx.shared import Pt
from unittest.mock import patch

from app.states.paper_data import (
    MAX_BYTES,
    extract_uploaded_paper,
    get_sample_papers,
)


class PaperDataTests(unittest.TestCase):
    def test_samples_are_complete_and_independent(self):
        papers = get_sample_papers()
        self.assertEqual(len(papers), 3)
        self.assertEqual(len({p["id"] for p in papers}), 3)
        for paper in papers:
            self.assertEqual(paper["source"], "sample")
            self.assertGreater(len(paper["passages"]), 3)
            self.assertTrue(paper["references"])
        papers[0]["passages"][0]["text"] = "changed"
        self.assertNotEqual(
            get_sample_papers()[0]["passages"][0]["text"], "changed"
        )

    def test_docx_extraction(self):
        doc = Document()
        doc.add_heading("Methods", level=1)
        doc.add_paragraph("We measured recall after a reading task.")
        doc.add_heading("Limitations", level=1)
        doc.add_paragraph("This was a small sample.")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Group"
        table.cell(0, 1).text = "Score"
        stream = BytesIO()
        doc.save(stream)
        paper = extract_uploaded_paper("reading.DOCX", stream.getvalue())
        self.assertEqual(paper["source"], "upload")
        self.assertEqual(paper["passages"][0]["section"], "Methods")
        self.assertIn("measured recall", paper["passages"][0]["text"])
        self.assertEqual(paper["passages"][-1]["section"], "Extracted table")
        self.assertTrue(all(p["page"] == 0 for p in paper["passages"]))
        self.assertEqual(paper["references"], [])

    def save_docx(self, document):
        stream = BytesIO()
        document.save(stream)
        return extract_uploaded_paper("sections.docx", stream.getvalue())

    def make_pdf(self, pages):
        with pymupdf.open() as document:
            for rows in pages:
                page = document.new_page()
                for index, row in enumerate(rows):
                    text, size, bold = row
                    page.insert_text(
                        (50, 50 + index * 25),
                        text,
                        fontsize=size,
                        fontname="hebo" if bold else "helv",
                    )
            return document.tobytes()

    def test_normal_docx_headings_and_wrapped_references(self):
        document = Document()
        for text in [
            "ABSTRACT",
            "A small study examines recall and attention.",
            "1 Introduction",
            "Prior work (Lee, 2021) motivates our question.",
            "METHODS",
            "We recorded scores after a reading task.",
            "REFERENCES",
            "Lee, A.",
            "(2021). Learning and attention.",
            "Journal of Learning, 12, 20-31.",
            "Unused, R. (2016). Separate research. Research Press.",
            "Appendix A: Materials",
            "Extra material cites (Other, 2022).",
        ]:
            document.add_paragraph(text)
        paper = self.save_docx(document)
        self.assertEqual(
            [p["section"] for p in paper["passages"]],
            ["ABSTRACT", "Introduction", "METHODS", "Appendix A: Materials"],
        )
        self.assertEqual(len(paper["references"]), 2)
        body = " ".join(p["text"] for p in paper["passages"])
        self.assertIn("(Lee, 2021)", body)
        self.assertNotIn("Journal of Learning", body)
        self.assertIn("(2021). Learning", paper["references"][0]["text"])
        self.assertTrue(paper["bibliography_detected"])

    def test_custom_docx_headings_and_table_order(self):
        document = Document()
        document.add_paragraph(
            "Unlabelled opening text remains available to the reader."
        )
        document.add_paragraph("2.1 Recruitment Strategy")
        document.add_paragraph(
            "We recruited volunteers from the local community."
        )
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Lee, A. (2021)"
        table.cell(0, 1).text = "20"
        document.add_paragraph(
            "After the table we describe the response rates."
        )
        paragraph = document.add_paragraph()
        paragraph.add_run("Unexpected Patterns").bold = True
        document.add_paragraph("The observations were mixed across groups.")
        paragraph = document.add_paragraph()
        paragraph.add_run("Practical implications").font.size = Pt(16)
        document.add_paragraph(
            "Larger studies are needed before adopting this approach."
        )
        paper = self.save_docx(document)
        sections = [p["section"] for p in paper["passages"]]
        self.assertEqual(
            sections,
            [
                "Document text",
                "Recruitment Strategy",
                "Extracted table",
                "Recruitment Strategy",
                "Unexpected Patterns",
                "Practical implications",
            ],
        )
        self.assertEqual(paper["references"], [])

    def test_pdf_mid_page_sections_and_multipage_bibliography(self):
        data = self.make_pdf(
            [
                [
                    ("Abstract", 12, False),
                    ("We examine recall in a small learning study.", 11, False),
                    ("1 Introduction", 12, False),
                    ("Prior work (Lee, 2021) motivates this study.", 11, False),
                    ("2 Recruitment Strategy", 12, False),
                    (
                        "We recruited adult volunteers from local classes.",
                        11,
                        False,
                    ),
                ],
                [
                    (
                        "Participants completed the same reading task.",
                        11,
                        False,
                    ),
                    ("Unexpected patterns", 15, True),
                    ("Scores differed across study groups.", 11, False),
                    ("Works Cited", 12, False),
                    (
                        "[1] Lee, A. (2021). Learning through explanation.",
                        11,
                        False,
                    ),
                    ("Journal of Learning, 12(2),", 11, False),
                ],
                [
                    ("20-31. https://doi.org/10.1234/test", 11, False),
                    (
                        "[2] Unused, R. (2016). A separate source. Research Press.",
                        11,
                        False,
                    ),
                    ("Supplementary material", 12, False),
                    (
                        "Extra results (Other, 2022) remain part of the body.",
                        11,
                        False,
                    ),
                ],
            ]
        )
        paper = extract_uploaded_paper("transitions.pdf", data)
        self.assertEqual(
            [(p["section"], p["page"]) for p in paper["passages"]],
            [
                ("Abstract", 1),
                ("Introduction", 1),
                ("Recruitment Strategy", 1),
                ("Recruitment Strategy", 2),
                ("Unexpected patterns", 2),
                ("Supplementary material", 3),
            ],
        )
        self.assertEqual(len(paper["references"]), 2)
        self.assertIn("20-31.", paper["references"][0]["text"])
        self.assertNotIn(
            "Research Press", " ".join(p["text"] for p in paper["passages"])
        )

    def test_bibliography_aliases_and_prose_are_not_entries(self):
        for heading in [
            "References",
            "Bibliography",
            "Works Cited",
            "Literature Cited",
            "References and Notes",
        ]:
            with self.subTest(heading=heading):
                document = Document()
                document.add_paragraph(
                    "The paper discusses Lee (2021) in ordinary prose."
                )
                document.add_paragraph(heading)
                document.add_paragraph(
                    "This sentence discusses Lee (2021) but is not a reference entry."
                )
                document.add_paragraph("1. We collected data in 2020.")
                document.add_paragraph(
                    "Lee, A. (2021). A source title. Research Press."
                )
                document.add_paragraph("2. Appendix A")
                document.add_paragraph(
                    "The appendix contains additional observations."
                )
                paper = self.save_docx(document)
                self.assertEqual(len(paper["references"]), 1)
                self.assertIn(
                    "This sentence",
                    " ".join(p["text"] for p in paper["passages"]),
                )
                self.assertEqual(paper["passages"][-1]["section"], "Appendix A")

    @patch("app.states.paper_data.pytesseract.image_to_string")
    def test_ocr_headings_use_same_parser(self, ocr):
        ocr.return_value = "ABSTRACT\nThis scanned study examines learning and recall.\nMETHODS\nWe recruited twenty volunteers.\nREFERENCES\nLee, A. (2021). Learning research. Research Press."
        data = self.make_pdf([[]])
        paper = extract_uploaded_paper("scan.pdf", data)
        self.assertEqual(
            [p["section"] for p in paper["passages"]], ["ABSTRACT", "METHODS"]
        )
        self.assertEqual(len(paper["references"]), 1)
        self.assertIn("English OCR was used", paper["note"])

    def test_pdf_without_headings_preserves_text(self):
        data = self.make_pdf(
            [
                [
                    (
                        "This ordinary sentence contains results (Lee, 2021).",
                        11,
                        False,
                    )
                ]
            ]
        )
        paper = extract_uploaded_paper("plain.pdf", data)
        self.assertEqual(paper["passages"][0]["section"], "Document text")
        self.assertIn("(Lee, 2021)", paper["passages"][0]["text"])
        self.assertEqual(paper["references"], [])

    def test_invalid_files(self):
        for filename, data in [
            ("notes.txt", b"hello"),
            ("paper.pdf", b"not a PDF"),
            ("paper.docx", b""),
            ("paper.pdf", b"x" * (MAX_BYTES + 1)),
        ]:
            with self.subTest(filename=filename, size=len(data)):
                with self.assertRaises(ValueError):
                    extract_uploaded_paper(filename, data)

    def test_blank_and_protected_pdfs(self):
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        stream = BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, "No readable text"):
            extract_uploaded_paper("blank.pdf", stream.getvalue())
        writer.encrypt("secret")
        stream = BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, "password-protected"):
            extract_uploaded_paper("protected.pdf", stream.getvalue())


if __name__ == "__main__":
    unittest.main()
