import reflex as rx

import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document

from app.citation_audit import (
    _build_explanation,
    _build_issues,
    _correction_prompt,
    _references_for_paper,
    failed_audit_report,
    run_citation_audit,
)
from app.states.paper_data import extract_uploaded_paper, get_sample_papers


def skipped_crossref(reference):
    """Stand in for a real lookup, stamping the same provenance fields."""
    from datetime import UTC, datetime

    reference["status"] = "lookup_error"
    reference["message"] = "Crossref disabled in unit test."
    reference["lookup_method"] = (
        "doi" if reference.get("doi") else "bibliographic_search"
    )
    reference["checked_at"] = datetime.now(UTC).isoformat(timespec="seconds")
    reference["uncertainty"] = "Provider was not contacted."
    return reference


class CitationAuditTests(unittest.TestCase):
    @patch("app.citation_audit._verify_reference", side_effect=skipped_crossref)
    def test_sample_findings_are_derived_from_paper_text(self, _verify):
        paper = get_sample_papers()[0]
        report = run_citation_audit(paper)
        categories = {issue["category"] for issue in report["issues"]}
        self.assertEqual(report["status"], "incomplete")
        self.assertIn("Citation missing from reference list", categories)
        self.assertIn("Possible duplicate reference", categories)
        self.assertIn("Citation placement · manual review", categories)
        self.assertTrue(all(issue["excerpt"] for issue in report["issues"]))
        self.assertIn(paper["title"], report["fix_prompt"])
        self.assertTrue(
            all(
                reference["status"] == "lookup_error"
                for reference in report["references"]
            )
        )

    @patch("app.citation_audit._verify_reference", side_effect=skipped_crossref)
    def test_docx_bibliography_is_parsed_and_compared(self, _verify):
        document = Document()
        document.add_heading("Introduction", level=1)
        document.add_paragraph(
            "Prior work supports this question (Lee et al., 2021). A second source is noted (Miller, 2018)."
        )
        document.add_heading("References", level=1)
        document.add_paragraph(
            "Lee, A. (2021). Learning through explanation. Journal of Learning, 12(2), 20–31."
        )
        document.add_paragraph(
            "Unused, R. (2016). A separate source. Research Press."
        )
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("audit.docx", stream.getvalue())
        report = run_citation_audit(paper)
        self.assertEqual(len(report["references"]), 2)
        self.assertTrue(
            any("References" in ref["location"] for ref in report["references"])
        )
        self.assertTrue(
            any(
                issue["category"] == "Citation missing from reference list"
                for issue in report["issues"]
            )
        )
        self.assertTrue(
            any(
                issue["category"] == "Uncited reference"
                for issue in report["issues"]
            )
        )
        self.assertTrue(all(issue["excerpt"] for issue in report["issues"]))

    def test_upload_references_are_authoritative_and_body_only_is_scanned(self):
        document = Document()
        document.add_paragraph("Introduction")
        document.add_paragraph("We cite Lee (2021) in the study text.")
        document.add_paragraph("References")
        document.add_paragraph("Lee, A. (2021). A study. Research Press.")
        document.add_paragraph(
            "Unused, R. (2016). Review of Other (2019). Research Press."
        )
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("body.docx", stream.getvalue())
        references, found = _references_for_paper(paper)
        self.assertTrue(found)
        self.assertEqual(
            [r["text"] for r in references],
            [r["text"] for r in paper["references"]],
        )
        issues = _build_issues(paper, references, found)
        self.assertFalse(
            any(
                i["category"] == "Citation missing from reference list"
                for i in issues
            )
        )
        self.assertEqual(
            sum(i["category"] == "Uncited reference" for i in issues), 1
        )
        paper["references"] = []
        self.assertEqual(_references_for_paper(paper), ([], True))

    def test_legacy_upload_bibliography_is_not_scanned_as_body(self):
        paper = get_sample_papers()[0]
        paper["source"] = "upload"
        paper["references"] = []
        paper["passages"] = [
            {
                "id": "p1",
                "section": "Extracted page",
                "page": 1,
                "text": "Introduction\nWe cite Lee (2021).\nReferences\nLee, A. (2021). A study. Research Press.\nUnused, R. (2016). Review of Other (2019). Research Press.\nAppendix A\nAdditional observations remain available.",
            }
        ]
        references, found = _references_for_paper(paper)
        self.assertTrue(found)
        self.assertEqual(len(references), 2)
        issues = _build_issues(paper, references, found)
        self.assertFalse(
            any(
                i["category"] == "Citation missing from reference list"
                for i in issues
            )
        )
        self.assertEqual(
            sum(i["category"] == "Uncited reference" for i in issues), 1
        )

    @patch("app.citation_audit._verify_reference", side_effect=skipped_crossref)
    def test_missing_bibliography_is_reported_as_incomplete(self, _verify):
        document = Document()
        document.add_heading("Findings", level=1)
        document.add_paragraph("An assertion appears here (Lopez, 2020).")
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("no_references.docx", stream.getvalue())
        report = run_citation_audit(paper)
        self.assertEqual(report["status"], "incomplete")
        self.assertTrue(
            any(
                issue["category"] == "Bibliography extraction"
                for issue in report["issues"]
            )
        )
        self.assertEqual(report["references"], [])

    @patch("app.citation_audit._verify_reference", side_effect=skipped_crossref)
    def test_numeric_citations_are_not_reported_as_missing_or_uncited(
        self, _verify
    ):
        document = Document()
        document.add_heading("Introduction", level=1)
        document.add_paragraph(
            "Earlier work established the baseline effect [1]. "
            "A later replication reported a smaller change [2, 3]."
        )
        document.add_heading("References", level=1)
        document.add_paragraph(
            "Lee, A. (2021). Learning through explanation. Journal of Learning, 12(2), 20–31."
        )
        document.add_paragraph("Miller, B. (2019). Replication notes. Research Press.")
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("numeric.docx", stream.getvalue())
        report = run_citation_audit(paper)
        categories = {issue["category"] for issue in report["issues"]}
        self.assertIn("Citation style not fully supported", categories)
        self.assertNotIn("Uncited reference", categories)
        self.assertNotIn("Citation missing from reference list", categories)
        self.assertEqual(report["status"], "incomplete")
        # The style limitation is reported as a structured reason, not only in
        # the prose summary.
        self.assertTrue(any("not fully supported" in l for l in report["limits"]))
        self.assertTrue(report["explanation"])
        self.assertTrue(all(issue["excerpt"] for issue in report["issues"]))

    def test_bracketed_years_are_not_mistaken_for_numeric_citations(self):
        document = Document()
        document.add_heading("Introduction", level=1)
        document.add_paragraph(
            "A review flagged the issue [2020] and agreed (Lee et al., 2021)."
        )
        document.add_heading("References", level=1)
        document.add_paragraph(
            "Lee, A. (2021). Learning through explanation. Journal of Learning, 12(2), 20–31."
        )
        document.add_paragraph("Miller, B. (2019). Replication notes. Research Press.")
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("bracketed.docx", stream.getvalue())
        references, found = _references_for_paper(paper)
        issues = _build_issues(paper, references, found)
        categories = {issue["category"] for issue in issues}
        self.assertNotIn("Citation style not fully supported", categories)
        self.assertEqual(
            sum(i["category"] == "Uncited reference" for i in issues), 1
        )


class AuditCompletionTests(unittest.TestCase):
    """Severity, structured explanation, limits, and prompt structure."""

    def audit(self, paper):
        with patch(
            "app.citation_audit._verify_reference", side_effect=skipped_crossref
        ):
            return run_citation_audit(paper)

    def test_every_finding_carries_a_severity(self):
        report = self.audit(get_sample_papers()[0])
        self.assertTrue(report["issues"])
        for issue in report["issues"]:
            with self.subTest(category=issue["category"]):
                self.assertIn(issue["severity"], {"high", "medium", "low"})

    def test_missing_reference_is_high_severity(self):
        report = self.audit(get_sample_papers()[0])
        severities = {
            issue["category"]: issue["severity"] for issue in report["issues"]
        }
        self.assertEqual(
            severities.get("Citation missing from reference list"), "high"
        )

    def test_report_explains_itself_in_sections(self):
        report = self.audit(get_sample_papers()[0])
        headings = [note["heading"] for note in report["explanation"]]
        for required in (
            "What was checked",
            "What was found",
            "What appears reliable",
            "What needs manual review",
            "What could not be checked",
            "What to do next",
        ):
            with self.subTest(heading=required):
                self.assertIn(required, headings)
        for note in report["explanation"]:
            self.assertTrue(note["body"].strip())

    def test_summary_cards_are_populated(self):
        report = self.audit(get_sample_papers()[0])
        counts = report["counts"]
        for key in (
            "citations_found",
            "references_parsed",
            "references_checked",
            "confirmed",
            "needing_review",
            "extraction_warnings",
        ):
            with self.subTest(key=key):
                self.assertIn(key, counts)
                self.assertIsInstance(counts[key], int)
        self.assertGreater(counts["references_parsed"], 0)
        self.assertGreater(counts["citations_found"], 0)

    def test_partial_audit_states_each_reason(self):
        report = self.audit(get_sample_papers()[0])
        self.assertEqual(report["status"], "incomplete")
        self.assertTrue(report["limits"])
        self.assertTrue(all(limit.strip() for limit in report["limits"]))

    def test_references_record_lookup_provenance(self):
        report = self.audit(get_sample_papers()[0])
        for reference in report["references"]:
            with self.subTest(reference=reference["id"]):
                for key in (
                    "parsed_author",
                    "parsed_year",
                    "parsed_title",
                    "lookup_method",
                    "matched_fields",
                    "conflicting_fields",
                    "uncertainty",
                    "checked_at",
                ):
                    self.assertIn(key, reference)
                self.assertTrue(reference["text"])
                # The original entry is never overwritten by parsed fields.
                self.assertNotEqual(reference["text"], reference["parsed_title"])

    def test_parsed_fields_come_from_the_entry(self):
        document = Document()
        document.add_heading("Introduction", level=1)
        document.add_paragraph("Prior work (Lee et al., 2021) supports this.")
        document.add_heading("References", level=1)
        document.add_paragraph(
            "Lee, A. B. (2021). Learning through explanation. "
            "Journal of Learning, 12(2), 20-31."
        )
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("parsed.docx", stream.getvalue())
        references, _ = _references_for_paper(paper)
        self.assertEqual(len(references), 1)
        reference = references[0]
        self.assertEqual(reference["parsed_year"], "2021")
        self.assertEqual(reference["parsed_author"], "Lee")
        self.assertIn("Learning through explanation", reference["parsed_title"])
        self.assertEqual(reference["parsed_venue"], "Journal of Learning")

    def test_venue_is_not_bleeding_into_the_title(self):
        from app.citation_audit import _reference_venue

        entry = (
            "Ward, A. F., Duke, K., Gneezy, A., & Bos, M. W. (2017). Brain "
            "drain: The mere presence of one's own smartphone reduces "
            "available cognitive capacity. Journal of the Association for "
            "Consumer Research, 2(2), 140-154."
        )
        venue = _reference_venue(entry)
        self.assertEqual(venue, "Journal of the Association for Consumer Research")
        self.assertNotIn("Brain drain", venue)

    def test_correction_prompt_is_structured_and_safe(self):
        report = self.audit(get_sample_papers()[0])
        prompt = report["fix_prompt"]
        for section in (
            "Paper under review",
            "Audit scope",
            "Important limitations",
            "Findings grouped by severity",
            "High severity",
            "Medium severity",
            "Low severity",
            "Reference verification notes",
            "Requested correction format",
            "Safety rules",
        ):
            with self.subTest(section=section):
                self.assertIn(section, prompt)
        self.assertIn(get_sample_papers()[0]["title"], prompt)
        self.assertIn("leads for review", prompt)
        self.assertIn("Do not invent", prompt)
        self.assertIn("citation style", prompt)

    def test_prompt_admits_a_clean_audit(self):
        paper = get_sample_papers()[0]
        report = self.audit(paper)
        # With no findings at all, the prompt must say so rather than imply
        # problems were found and suppressed.
        clean = _correction_prompt(paper, [], report["references"], [])
        self.assertIn("None.", clean)
        self.assertIn("leads for review", clean)
        notes = _build_explanation(
            report["counts"], [], report["references"], []
        )
        next_steps = next(
            note["body"] for note in notes if note["heading"] == "What to do next"
        )
        self.assertIn("spot-check", next_steps.lower())
        review = next(
            note["body"]
            for note in notes
            if note["heading"] == "What needs manual review"
        )
        self.assertIn("Nothing was flagged", review)

    def test_numeric_citation_numbers_out_of_range_are_flagged(self):
        document = Document()
        document.add_heading("Introduction", level=1)
        document.add_paragraph("Earlier work showed an effect [7] and later work agreed [9].")
        document.add_heading("References", level=1)
        document.add_paragraph("Lee, A. (2021). A study. Research Press.")
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("numericrange.docx", stream.getvalue())
        report = self.audit(paper)
        categories = {issue["category"] for issue in report["issues"]}
        self.assertIn("Citation number out of range", categories)
        finding = next(
            issue
            for issue in report["issues"]
            if issue["category"] == "Citation number out of range"
        )
        self.assertEqual(finding["severity"], "medium")
        self.assertIn("[9]", finding["explanation"])

    def test_in_range_numeric_citations_are_not_flagged(self):
        document = Document()
        document.add_heading("Introduction", level=1)
        document.add_paragraph("Earlier work showed an effect [1] and later work agreed [2].")
        document.add_heading("References", level=1)
        document.add_paragraph("Lee, A. (2021). A study. Research Press.")
        document.add_paragraph("Miller, B. (2019). Another study. Research Press.")
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("numericok.docx", stream.getvalue())
        report = self.audit(paper)
        categories = {issue["category"] for issue in report["issues"]}
        self.assertNotIn("Citation number out of range", categories)

    def test_failed_report_still_has_the_new_shape(self):
        report = failed_audit_report("x", "boom")
        self.assertEqual(report["status"], "error")
        self.assertIn("counts", report)
        self.assertEqual(report["explanation"], [])
        self.assertEqual(report["limits"], [])


if __name__ == "__main__":
    unittest.main()
