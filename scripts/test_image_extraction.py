import os
import shutil
import tempfile
import unittest
import unittest.mock
from io import BytesIO

from docx import Document
from docx.shared import Inches
from PIL import Image as PILImage
import pymupdf
import pytesseract

from app import image_extraction
from app.image_extraction import (
    CAPTION_PATTERN,
    MAX_IMAGES,
    _classify_image,
    _is_caption,
    discard_media,
    extract_docx_images,
    extract_pdf_images,
    media_dir,
)
from app.states.paper_data import extract_uploaded_paper, get_sample_papers


VALID_PAPER_ID = "upload-" + "a" * 32
OTHER_PAPER_ID = "upload-" + "b" * 32


def figure_png(width: int = 600, height: int = 400, colour=(20, 90, 160)) -> bytes:
    image = PILImage.new("RGB", (width, height), "white")
    for y in range(0, height, 40):
        for dx in range(6):
            image.putpixel((min(dx, width - 1), y), colour)
    stream = BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


class ImageExtractionTests(unittest.TestCase):
    def setUp(self):
        self.media = tempfile.mkdtemp(prefix="paperproof-test-media-")
        self._previous_root = os.environ.get("PAPERPROOF_MEDIA_DIR")
        os.environ["PAPERPROOF_MEDIA_DIR"] = self.media

    def tearDown(self):
        shutil.rmtree(self.media, ignore_errors=True)
        if self._previous_root is None:
            os.environ.pop("PAPERPROOF_MEDIA_DIR", None)
        else:
            os.environ["PAPERPROOF_MEDIA_DIR"] = self._previous_root

    # ---------------------------------------------------------------- captions

    def test_caption_pattern_recognises_common_forms(self):
        for text in (
            "Figure 1. Mean recall by condition",
            "Fig. 2: Daily ratings",
            "Table 3 - Sample characteristics",
            "Tab. 4",
            "Figure A1. Overview",
        ):
            with self.subTest(text=text):
                self.assertTrue(_is_caption(text))
        for text in (
            "The figure is shown below.",
            "Results",
            "We tabulated the values in Table 1 earlier.",
            "",
        ):
            with self.subTest(text=text):
                self.assertFalse(_is_caption(text))
        self.assertIsNotNone(CAPTION_PATTERN)

    def test_image_classification_uses_caption_wording(self):
        cases = {
            "Table 1. Participant characteristics": "table_image",
            "Figure 1. The reaction equation used": "equation_image",
            "Figure 2. Bar chart of mean scores": "chart",
            "Figure 3. Study workflow schematic": "diagram",
            "Figure 4. Micrograph of the stained sample": "photograph",
            "Figure 5. Mean recall by condition": "figure",
        }
        for caption, expected in cases.items():
            with self.subTest(caption=caption):
                self.assertEqual(_classify_image(caption, "", 0), expected)
        # No caption and no text: stay honest rather than guessing a kind.
        self.assertEqual(_classify_image("", "", 0), "unknown")

    # ---------------------------------------------------------------- storage

    def test_media_dir_rejects_untrusted_paper_ids(self):
        self.assertEqual(
            media_dir(VALID_PAPER_ID), self.media_path(VALID_PAPER_ID)
        )
        for bad in (
            "sample-attention",
            "../../etc",
            "upload-../../windows",
            "upload-nothex",
            "",
        ):
            with self.subTest(paper_id=bad):
                with self.assertRaises(ValueError):
                    media_dir(bad)

    def media_path(self, paper_id):
        from pathlib import Path

        return Path(self.media) / paper_id

    def test_discard_media_removes_stored_images(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_image(pymupdf.Rect(50, 50, 400, 350), stream=figure_png())
        data = document.tobytes()
        document.close()
        images, _ = extract_pdf_images(VALID_PAPER_ID, data, [])
        self.assertEqual(len(images), 1)
        self.assertTrue(os.path.exists(images[0]["local_path"]))
        discard_media(VALID_PAPER_ID)
        self.assertFalse(os.path.exists(images[0]["local_path"]))
        # Discarding twice must not raise.
        discard_media(VALID_PAPER_ID)

    # ---------------------------------------------------------------- PDF

    def build_pdf(self):
        document = pymupdf.open()
        first = document.new_page()
        first.insert_textbox(
            pymupdf.Rect(50, 50, 550, 80), "Abstract", fontsize=15
        )
        first.insert_textbox(
            pymupdf.Rect(50, 90, 550, 115),
            "Figure 1. Mean recall score by condition.",
            fontsize=10,
        )
        first.insert_image(
            pymupdf.Rect(120, 130, 480, 340), stream=figure_png()
        )
        first.insert_textbox(
            pymupdf.Rect(50, 355, 550, 430),
            "Eighty students completed a reading task with the phone on the "
            "desk or in another room.",
            fontsize=10,
        )
        second = document.new_page()
        second.insert_textbox(
            pymupdf.Rect(50, 50, 550, 80), "Methods", fontsize=15
        )
        second.insert_textbox(
            pymupdf.Rect(50, 90, 550, 115),
            "Table 1. Participant characteristics.",
            fontsize=10,
        )
        second.insert_image(
            pymupdf.Rect(120, 130, 480, 340),
            stream=figure_png(colour=(200, 40, 40)),
        )
        second.insert_textbox(
            pymupdf.Rect(50, 355, 550, 430),
            "Participants were recruited from two introductory classes.",
            fontsize=10,
        )
        data = document.tobytes()
        document.close()
        return data

    def test_pdf_images_are_extracted_in_document_order(self):
        images, _ = extract_pdf_images(
            VALID_PAPER_ID, self.build_pdf(), []
        )
        self.assertEqual(len(images), 2)
        self.assertEqual([i["order"] for i in images], [0, 1])
        self.assertEqual([i["id"] for i in images], ["img-01", "img-02"])
        self.assertEqual([i["page"] for i in images], [1, 2])
        for image in images:
            with self.subTest(image=image["id"]):
                self.assertTrue(os.path.exists(image["local_path"]))
                self.assertEqual(image["width"], 600)
                self.assertEqual(image["height"], 400)
                self.assertEqual(image["mime_type"], "image/png")
                self.assertEqual(image["extraction_status"], "extracted")
                self.assertEqual(image["analysis_status"], "not_run")

    def test_pdf_caption_and_section_are_associated(self):
        images, _ = extract_pdf_images(
            VALID_PAPER_ID, self.build_pdf(), []
        )
        first, second = images
        self.assertEqual(first["caption"], "Figure 1. Mean recall score by condition.")
        self.assertEqual(first["section"], "Abstract")
        self.assertIn("Eighty students", first["surrounding_text"])
        self.assertEqual(second["caption"], "Table 1. Participant characteristics.")
        self.assertEqual(second["section"], "Methods")
        self.assertEqual(second["image_kind"], "table_image")

    def test_pdf_without_figures_returns_empty_and_warns(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_textbox(
            pymupdf.Rect(50, 50, 550, 200),
            "Abstract\nA short paper with no figures at all.",
            fontsize=11,
        )
        data = document.tobytes()
        document.close()
        images, _ = extract_pdf_images(VALID_PAPER_ID, data, [])
        self.assertEqual(images, [])

    def test_tiny_images_are_skipped(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_textbox(
            pymupdf.Rect(50, 50, 550, 120),
            "Abstract\nA decorative bullet glyph follows.",
            fontsize=11,
        )
        page.insert_image(
            pymupdf.Rect(50, 130, 62, 142), stream=figure_png(12, 12)
        )
        data = document.tobytes()
        document.close()
        images, _ = extract_pdf_images(VALID_PAPER_ID, data, [])
        self.assertEqual(images, [])

    def test_image_limit_is_declared(self):
        self.assertEqual(MAX_IMAGES, 200)

    # ---------------------------------------------------------------- DOCX

    def build_docx(self):
        document = Document()
        document.add_heading("Abstract", level=1)
        document.add_paragraph("Sixty students recorded outdoor minutes daily.")
        document.add_heading("Methods", level=1)
        document.add_paragraph("Wellbeing was rated from 1 to 7 for fourteen days.")
        document.add_paragraph("Figure 2. Daily wellbeing ratings by condition.")
        document.add_picture(BytesIO(figure_png(500, 350)), width=Inches(4))
        document.add_heading("Results", level=1)
        document.add_paragraph("Higher-exposure days scored higher overall.")
        document.add_paragraph("Figure 3. Distribution of daily scores.")
        document.add_picture(BytesIO(figure_png(480, 320)), width=Inches(4))
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    def test_docx_images_keep_document_order_and_captions(self):
        images, _ = extract_docx_images(
            VALID_PAPER_ID,
            Document(BytesIO(self.build_docx())),
            [],
        )
        self.assertEqual(len(images), 2)
        self.assertEqual([i["order"] for i in images], [0, 1])
        self.assertEqual([i["id"] for i in images], ["img-01", "img-02"])
        first, second = images
        self.assertEqual(
            first["caption"], "Figure 2. Daily wellbeing ratings by condition."
        )
        self.assertEqual(first["section"], "Methods")
        self.assertEqual(
            second["caption"], "Figure 3. Distribution of daily scores."
        )
        self.assertEqual(second["section"], "Results")
        for image in images:
            with self.subTest(image=image["id"]):
                self.assertTrue(os.path.exists(image["local_path"]))
                self.assertEqual(image["mime_type"], "image/png")
                self.assertEqual(image["width"], 500 if image is first else 480)

    def test_docx_images_report_no_page_number(self):
        images, warnings = extract_docx_images(
            VALID_PAPER_ID, Document(BytesIO(self.build_docx())), []
        )
        self.assertTrue(images)
        for image in images:
            self.assertEqual(image["page"], 0)
        codes = {item["code"] for item in warnings}
        self.assertIn("image_kind_inferred", codes)

    # ------------------------------------------------- paper integration

    def test_upload_without_figures_records_an_explicit_warning(self):
        document = Document()
        document.add_heading("Abstract", level=1)
        document.add_paragraph("A short paper with no figures at all.")
        stream = BytesIO()
        document.save(stream)
        paper = extract_uploaded_paper("plain.docx", stream.getvalue())
        self.assertEqual(paper["images"], [])
        codes = {item["code"] for item in paper["extraction_warnings"]}
        self.assertIn("no_images_found", codes)
        self.assertIn("docx_page_numbers_unavailable", codes)

    def test_uploaded_pdf_exposes_images_and_warnings(self):
        paper = extract_uploaded_paper("figures.pdf", self.build_pdf())
        self.assertEqual(len(paper["images"]), 2)
        self.assertTrue(paper["extraction_warnings"])
        for warning in paper["extraction_warnings"]:
            with self.subTest(code=warning["code"]):
                self.assertIn(
                    warning["severity"], {"info", "warning", "error"}
                )
                self.assertTrue(warning["message"])
                self.assertTrue(warning["recommended_action"])
        self.assertIn(
            "research_classification", paper
        )
        self.assertTrue(paper["research_classification"]["outcome"])

    def test_uploaded_docx_exposes_images(self):
        paper = extract_uploaded_paper("figures.docx", self.build_docx())
        self.assertEqual(len(paper["images"]), 2)
        self.assertTrue(paper["extraction_warnings"])

    def test_sample_papers_have_no_images_and_are_independent(self):
        papers = get_sample_papers()
        for paper in papers:
            with self.subTest(paper=paper["id"]):
                self.assertEqual(paper["images"], [])
                self.assertEqual(paper["extraction_warnings"], [])
                self.assertEqual(
                    paper["research_classification"]["outcome"], "accepted"
                )
        papers[0]["images"].append({"id": "x"})
        self.assertEqual(get_sample_papers()[0]["images"], [])

    def test_each_paper_owns_a_separate_media_directory(self):
        images, _ = extract_pdf_images(VALID_PAPER_ID, self.build_pdf(), [])
        other, _ = extract_pdf_images(OTHER_PAPER_ID, self.build_pdf(), [])
        self.assertNotEqual(images[0]["local_path"], other[0]["local_path"])
        discard_media(VALID_PAPER_ID)
        self.assertTrue(os.path.exists(other[0]["local_path"]))
        discard_media(OTHER_PAPER_ID)


class OcrAvailabilityTests(unittest.TestCase):
    """Local OCR must be reported honestly, before a paper is opened.

    These tests fail without `ocr_availability`: the probe does not exist, so
    a machine without Tesseract has no way to learn that figure text is being
    silently dropped.
    """

    def setUp(self):
        self._probe = image_extraction._ocr_probe

    def tearDown(self):
        image_extraction._ocr_probe = self._probe

    def test_missing_tesseract_is_reported_as_unavailable(self):
        image_extraction._ocr_probe = None
        with unittest.mock.patch.object(
            image_extraction.pytesseract,
            "get_tesseract_version",
            side_effect=RuntimeError("not installed"),
        ):
            result = image_extraction.ocr_availability()
        self.assertFalse(result["available"])
        self.assertEqual(result["version"], "")
        # The message must tell the reader what changes for them, not just that
        # something is broken.
        self.assertIn("captions", result["detail"])
        self.assertIn("Tesseract", result["detail"])

    def test_missing_english_pack_is_reported_as_unavailable(self):
        image_extraction._ocr_probe = None
        with unittest.mock.patch.object(
            image_extraction.pytesseract,
            "get_tesseract_version",
            return_value="5.5.3",
        ), unittest.mock.patch.object(
            image_extraction.pytesseract,
            "get_languages",
            return_value=["osd"],
        ):
            result = image_extraction.ocr_availability()
        self.assertFalse(result["available"])
        self.assertIn("5.5.3", result["version"])
        self.assertIn("English", result["detail"])

    def test_working_tesseract_reports_available(self):
        image_extraction._ocr_probe = None
        with unittest.mock.patch.object(
            image_extraction.pytesseract,
            "get_tesseract_version",
            return_value="5.5.3",
        ), unittest.mock.patch.object(
            image_extraction.pytesseract,
            "get_languages",
            return_value=["eng", "osd"],
        ):
            result = image_extraction.ocr_availability()
        self.assertTrue(result["available"])
        self.assertEqual(result["version"], "5.5.3")

    def test_probe_runs_once_per_process(self):
        image_extraction._ocr_probe = None
        with unittest.mock.patch.object(
            image_extraction.pytesseract,
            "get_tesseract_version",
            return_value="5.5.3",
        ) as version:
            first = image_extraction.ocr_availability()
            second = image_extraction.ocr_availability()
        # A computed var re-reads this on every render, so an uncached probe
        # would spawn a Tesseract subprocess per frame.
        self.assertEqual(version.call_count, 1)
        self.assertEqual(first, second)

    def test_probe_honours_force(self):
        image_extraction._ocr_probe = None
        with unittest.mock.patch.object(
            image_extraction.pytesseract,
            "get_tesseract_version",
            return_value="5.5.3",
        ) as version:
            image_extraction.ocr_availability()
            image_extraction.ocr_availability(force=True)
        self.assertEqual(version.call_count, 2)

    def test_probe_failure_is_logged_by_type_not_by_message(self):
        image_extraction._ocr_probe = None
        secret = "C:\\Users\\someone\\AppData secret-token"
        with unittest.mock.patch.object(
            image_extraction.pytesseract,
            "get_tesseract_version",
            side_effect=RuntimeError(secret),
        ), self.assertLogs(level="WARNING") as logs:
            image_extraction.ocr_availability()
        logged = "\n".join(logs.output)
        self.assertNotIn("secret-token", logged)
        self.assertIn("RuntimeError", logged)


class RealOcrProbeTests(unittest.TestCase):
    def setUp(self):
        self._probe = image_extraction._ocr_probe

    def tearDown(self):
        image_extraction._ocr_probe = self._probe

    def test_probe_agrees_with_a_real_tesseract_run(self):
        """The reported verdict must match what OCR actually does here."""
        image_extraction._ocr_probe = None
        result = image_extraction.ocr_availability()
        if not result["available"]:
            self.skipTest("Tesseract OCR is not available on this machine")
        text = render_ocr_probe_image()
        self.assertTrue(text.strip(), "OCR reported available but read nothing")


def render_ocr_probe_image() -> str:
    """Draw a known word and OCR it, the same way figures are read."""
    from PIL import ImageDraw

    with tempfile.TemporaryDirectory(prefix="paperproof-ocr-") as folder:
        path = os.path.join(folder, "probe.png")
        image = PILImage.new("RGB", (900, 200), "white")
        ImageDraw.Draw(image).text(
            (30, 70), "REFERENCE", fill="black", font_size=60
        )
        image.save(path)
        try:
            return pytesseract.image_to_string(
                PILImage.open(path), lang="eng", config="--psm 6"
            )
        except Exception as e:
            raise AssertionError(
                f"OCR reported available but the call failed: {type(e).__name__}"
            ) from e


if __name__ == "__main__":
    unittest.main()